# -*- coding: utf-8 -*-
"""Fill only uniquely matched, empty curated setlists from the Google sheet.

The sheet is authoritative for its setlists, but already published nonempty
records are reviewed by people. This job never creates an event or changes
cast, dates, links, or an existing song table.
"""

import argparse
import datetime as dt
import html
import io
import json
import os
from pathlib import Path
import re
import tempfile
import urllib.request


ROOT = Path(__file__).resolve().parent.parent
LIVE_CAT = ROOT / "data" / "live_cat_data.json"
REPORT = ROOT / "data" / "events" / "sheet_apply_report.json"
SHEET_URL = (
    "https://docs.google.com/spreadsheets/d/"
    "1HU5lkpjmz3_tXtJnHLvBcbeUXe_2LliD9AagFdwhDvk/export?format=xlsx"
)
SONG_RE = re.compile(r"^M\s*(\d+)[.\s]\s*(.*)$", re.I)
TITLE_SPLIT_RE = re.compile(r"[＠@]", re.I)
REGION_RE = re.compile(r"\s*[（(][^()（）]*[）)]\s*$")
DATE_RE = re.compile(r"\d{4}-\d{2}-\d{2}")
PERFORMER_SPLIT_RE = re.compile(r"[、，,､/・]")


def normalized_title(value):
    title = TITLE_SPLIT_RE.split(str(value or "").strip(), 1)[0]
    title = REGION_RE.sub("", title).strip()
    return re.sub(r"\s+", "", title).casefold()


def event_date(value):
    if isinstance(value, (dt.date, dt.datetime)):
        return value.isoformat()[:10]
    match = DATE_RE.search(str(value or ""))
    return match.group(0) if match else ""


def session_label(title):
    match = re.search(r"(?:DAY\s*\d+|[1-3]部|第[一二三四]部|アフターパート)", title or "", re.I)
    return match.group(0) if match else "本公演"


def split_performers(value):
    return [item.strip() for item in PERFORMER_SPLIT_RE.split(str(value)) if item.strip()]


def parse_sheet(data):
    import openpyxl

    workbook = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        sheet = workbook[workbook.sheetnames[0]]
        blocks = []
        current = None
        current_date = ""
        for row in sheet.iter_rows(values_only=True):
            date, title, song, performers = (list(row[:4]) + [None] * 4)[:4]
            if date is not None:
                current_date = event_date(date)
            if title is not None and str(title).strip():
                current = {"date": current_date, "title": str(title).strip(), "songs": []}
                blocks.append(current)
            elif song is not None and str(song).strip() and current is not None:
                text = str(song).strip()
                match = SONG_RE.match(text)
                current["songs"].append({
                    "no": (match.group(1).lstrip("0") or "0") if match else "",
                    "name": match.group(2).strip() if match else text,
                    "performers": split_performers(performers) if performers else [],
                })
            elif performers is not None and str(performers).strip() and current and current["songs"]:
                if not current["songs"][-1]["performers"]:
                    current["songs"][-1]["performers"] = split_performers(performers)
        return [block for block in blocks if block["songs"]]
    finally:
        workbook.close()


def curated_slots(catalog):
    slots = []
    for category in ("cd", "other", "twinkle"):
        root = catalog.get(category) or {}
        for section in root.get("sections") or [root]:
            for group in section.get("groups") or []:
                for sub in group.get("subs") or []:
                    for day in sub.get("days") or []:
                        slots.append({
                            "group": group.get("group") or "",
                            "title": sub.get("title") or "",
                            "date": event_date(sub.get("date")),
                            "label": day.get("label") or "本公演",
                            "cast": sub.get("cast") or "",
                            "day": day,
                        })
    return slots


def confirmed_cast(cast):
    """Only explicitly credited Uma performers may receive character avatars."""
    out = {}
    for actor, character in re.findall(r"([^、（）<>]+)（([^（）<>]+)）", cast or ""):
        out[actor.strip()] = character.strip()
    return out


def song_table(block, cast):
    from_map = confirmed_cast(cast)
    rows = []
    for index, song in enumerate(block["songs"], 1):
        performers = []
        for name in song["performers"]:
            safe = html.escape(name, quote=True)
            if name in from_map:
                # The canonical character identity is derived from the verified
                # event cast, never from a global actor-to-character lookup.
                character = html.escape(from_map[name], quote=True)
                performers.append('<span class="perf-item"><span class="perf-name">%s</span></span>' % character)
            else:
                performers.append('<span class="guest-performer">%s</span>' % safe)
        rows.append(
            '<tr><td class="setlist-no">%s</td><td class="setlist-song">%s</td><td class="setlist-perf">%s</td></tr>'
            % (html.escape(song["no"] or str(index)), html.escape(song["name"]), "".join(performers))
        )
    return (
        '<table class="setlist-table">\n'
        '<thead><tr><th>#</th><th>曲名</th><th>出演者</th></tr></thead>\n'
        '<tbody>\n%s\n</tbody>\n</table>' % "\n".join(rows)
    )


def table_summary(value):
    return [html.unescape(re.sub(r"<[^>]+>", "", item)).strip()
            for item in re.findall(r'<td class="setlist-song">([\s\S]*?)</td>', value or "")]


def table_credits(value, cast):
    by_character = {character: actor for actor, character in confirmed_cast(cast).items()}
    credits = []
    for row in re.findall(r"<tr\b[^>]*>([\s\S]*?)</tr>", value or "", re.I):
        if 'class="setlist-song"' not in row:
            continue
        names = [html.unescape(name).strip() for name in
                 re.findall(r'<span class="(?:perf-name|guest-performer)">([\s\S]*?)</span>', row)]
        credits.append([by_character.get(name, name) for name in names])
    return credits


def compare(blocks, catalog):
    slots = curated_slots(catalog)
    report = {"sheet_blocks": len(blocks), "filled": [], "review": [], "unchanged": 0}
    for block in blocks:
        title = normalized_title(block["title"])
        date = block["date"]
        label = session_label(block["title"])
        # No fuzzy guesses: a single title/date/session match is required.
        matches = [slot for slot in slots if normalized_title(slot["title"]) == title
                   and slot["date"] == date and slot["label"] == label]
        descriptor = {"sheet_title": block["title"], "date": date, "session": label}
        if len(matches) != 1:
            report["review"].append({**descriptor, "reason": "unmatched_or_ambiguous", "matches": len(matches)})
            continue
        slot = matches[0]
        existing = slot["day"].get("table") or ""
        sheet_songs = [song["name"] for song in block["songs"]]
        if existing.strip():
            published_credits = table_credits(existing, slot["cast"])
            sheet_credits = [song["performers"] for song in block["songs"]]
            same_credits = len(published_credits) == len(sheet_credits) and all(
                sorted(published) == sorted(source)
                for published, source in zip(published_credits, sheet_credits)
            )
            if table_summary(existing) == sheet_songs and same_credits:
                report["unchanged"] += 1
            else:
                report["review"].append({**descriptor, "reason": "nonempty_setlist_differs",
                                         "published_songs": table_summary(existing), "sheet_songs": sheet_songs,
                                         "published_performers": published_credits,
                                         "sheet_performers": sheet_credits})
            continue
        if not slot["cast"].strip() or not confirmed_cast(slot["cast"]):
            report["review"].append({**descriptor, "reason": "cast_not_confirmed"})
            continue
        if any(not song["performers"] for song in block["songs"]):
            report["review"].append({**descriptor, "reason": "missing_song_performers"})
            continue
        slot["day"]["table"] = song_table(block, slot["cast"])
        report["filled"].append({**descriptor, "published_title": slot["title"], "songs": len(sheet_songs)})
    return report


def write_json(path, value):
    path.parent.mkdir(parents=True, exist_ok=True)
    fd, temporary = tempfile.mkstemp(prefix=path.name + ".", suffix=".tmp", dir=str(path.parent))
    try:
        with os.fdopen(fd, "w", encoding="utf-8", newline="\n") as stream:
            json.dump(value, stream, ensure_ascii=False, indent=2)
            stream.write("\n")
        os.replace(temporary, path)
    finally:
        if os.path.exists(temporary):
            os.unlink(temporary)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--sheet-xlsx", type=Path, help="Use a local workbook instead of downloading")
    parser.add_argument("--live-cat", type=Path, default=LIVE_CAT)
    parser.add_argument("--report", type=Path, default=REPORT)
    parser.add_argument("--dry-run", action="store_true")
    args = parser.parse_args()
    if args.sheet_xlsx:
        data = args.sheet_xlsx.read_bytes()
    else:
        request = urllib.request.Request(SHEET_URL, headers={"User-Agent": "uma-live-wiki sheet-sync"})
        with urllib.request.urlopen(request, timeout=45) as response:
            data = response.read()
    if not data.startswith(b"PK"):
        raise ValueError("Google sheet response is not an xlsx workbook")
    blocks = parse_sheet(data)
    if not blocks:
        raise ValueError("Google sheet has no setlists; refusing to update curated data")
    catalog = json.loads(args.live_cat.read_text(encoding="utf-8"))
    report = compare(blocks, catalog)
    if report["filled"] and not args.dry_run:
        write_json(args.live_cat, catalog)
    write_json(args.report, report)
    print("sheet-sync: %d blocks, %d filled, %d unchanged, %d require review" %
          (report["sheet_blocks"], len(report["filled"]), report["unchanged"], len(report["review"])))


if __name__ == "__main__":
    main()
