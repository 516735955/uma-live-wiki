# -*- coding: utf-8 -*-
"""Crawl Umamusume events from Eventernote -> events_data.json
Usage: python3 crawl_events.py
"""
import io, sys, json, re, html, time, os, urllib.request
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

TOOLS_DIR = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(TOOLS_DIR)
OUT = os.path.join(ROOT, 'events_data.json')
BASE = 'https://www.eventernote.com/actors/%E3%82%A6%E3%83%9E%E5%A8%98%20%E3%83%97%E3%83%AA%E3%83%86%E3%82%A3%E3%83%BC%E3%83%80%E3%83%BC%E3%83%93%E3%83%BC(gal%E2%80%99up!)/22274/events'
LIMIT = 100
UA = {'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64)', 'Accept-Language': 'ja'}

DATE_RE = re.compile(r'<p class="day\d">\s*(\d{4}-\d{2}-\d{2})\s*\(\s*<span class="wday\d">([^<]+)</span>\s*\)', re.S)
IMG_RE = re.compile(r'<img src="([^"]+)"[^>]*>', re.S)
H4_RE = re.compile(r'<h4>\s*<a href="([^"]+)">(.*?)</a>', re.S)
VENUE_RE = re.compile(r'会場:\s*<a href="([^"]+)">([^<]+)</a>', re.S)
TIMES_RE = re.compile(r'<span class="s">([^<]*)</p>', re.S)
ACTOR_RE = re.compile(r'<li><a href="([^"]+)">([^<]+)</a></li>', re.S)


def fetch(url):
    return urllib.request.urlopen(urllib.request.Request(url, headers=UA), timeout=40).read().decode('utf-8', 'replace')


def extract_blocks(content):
    blocks = []
    starts = [m.start() for m in re.finditer(r'<li class="clearfix', content)]
    for i, s in enumerate(starts):
        depth = 0
        pos = s
        n = len(content)
        while pos < n:
            nxt_open = content.find('<li', pos)
            nxt_close = content.find('</li>', pos)
            if nxt_close == -1:
                break
            if nxt_open != -1 and nxt_open < nxt_close:
                depth += 1
                pos = nxt_open + 3
            else:
                depth -= 1
                pos = nxt_close + 5
                if depth == 0:
                    blocks.append(content[s:pos])
                    break
    return blocks


def parse_block(block):
    dm = DATE_RE.search(block)
    if not dm:
        return None
    img_m = IMG_RE.search(block)
    h4 = H4_RE.search(block)
    ven = VENUE_RE.search(block)
    tm = TIMES_RE.search(block)
    actors = []
    seen = set()
    for am in ACTOR_RE.finditer(block):
        name = html.unescape(am.group(2)).strip()
        if name in seen:
            continue
        seen.add(name)
        actors.append({'name': name, 'url': 'https://www.eventernote.com' + am.group(1)})
    return {
        'date': dm.group(1),
        'date_zh': dm.group(1) + ' (' + dm.group(2).strip() + ')',
        'title': html.unescape(h4.group(2)).strip() if h4 else '',
        'link': 'https://www.eventernote.com' + h4.group(1) if h4 else '',
        'img': img_m.group(1) if img_m else '',
        'venue': html.unescape(ven.group(2)).strip() if ven else '',
        'venue_link': 'https://www.eventernote.com' + ven.group(1) if ven else '',
        'times': tm.group(1).strip() if tm else '',
        'actors': actors,
        'live': '',
    }


LIVE_DATA = os.path.join(ROOT, 'live_cat_data.json')
NUMBERED_LIVE_DATA = os.path.join(ROOT, 'live_data.json')
APP_JS = os.path.join(ROOT, 'uma_tools', 'app.js')

EVENTS_XLSX = os.path.join(ROOT, 'events_list.xlsx')
VOICE_XLSX = os.path.join(ROOT, 'voice_list.xlsx')


def export_events_xlsx(new_events):
    """Insert freshly crawled events into events_list.xlsx at row 2 (shifting the
    rest down). 出演者 is filtered against voice_list.xlsx 日文名 column (col B);
    names not present there are dropped. Exception: '佐伯伊織(NU-KO)' is rewritten
    to '佐伯伊織' before matching."""
    if not new_events:
        print('export xlsx: no new events', flush=True)
        return
    # 1) load voice_list 日文名 (column B)
    vnames = set()
    try:
        import openpyxl
        vwb = openpyxl.load_workbook(VOICE_XLSX, data_only=True)
        vws = vwb[vwb.sheetnames[0]]
        for r in range(2, vws.max_row + 1):
            v = vws.cell(r, 2).value
            if v:
                vnames.add(str(v).strip())
    except Exception as ex:
        print('export xlsx: voice_list load failed, skip:', str(ex)[:80], flush=True)
        return
    # 2) build rows: 日期 / 标题 / 会场 / 出演者
    rows = []
    for e in new_events:
        date = e.get('date_zh') or e.get('date', '')
        title = e.get('title', '')
        venue = e.get('venue', '')
        kept = []
        for a in e.get('actors', []):
            nm = (a.get('name') or '').strip()
            if not nm:
                continue
            cand = '佐伯伊織' if nm == '佐伯伊織(NU-KO)' else nm
            if cand in vnames:
                kept.append(cand)
        rows.append([date, title, venue, '、'.join(kept)])
    # 3) insert at row 2 (header stays at row 1, former row 2 shifts down)
    try:
        import openpyxl
        wb = openpyxl.load_workbook(EVENTS_XLSX)
        ws = wb[wb.sheetnames[0]]
        n = len(rows)
        ws.insert_rows(2, n)
        for i, row in enumerate(rows):
            for c, val in enumerate(row, start=1):
                ws.cell(2 + i, c, val)
        wb.save(EVENTS_XLSX)
        print('export xlsx: inserted %d new rows into %s' % (n, EVENTS_XLSX), flush=True)
    except Exception as ex:
        print('export xlsx: write failed:', str(ex)[:80], flush=True)

_TRANS = str.maketrans({
    '　': ' ', ' ': '', '～': '~', '〜': '~', '－': '-', '―': '-', '−': '-',
    '！': '!', '？': '?', '＆': '&', '’': "'", '”': '"', '“': '"',
    '『': '', '』': '', '「': '', '」': '', '【': '', '】': '', '（': '', '）': '', '(': '', ')': '',
})
import re as _re
_DAY_RE = _re.compile(r'(\d+)\s*日目')
_YEAR_RE = _re.compile(r'(20\d{2})')


def _norm(s):
    s = (s or '').replace('配信', '')
    s = s.translate(_TRANS)
    s = _DAY_RE.sub(r'DAY\1', s)
    s = s.upper()
    s = s.replace('ウマ娘プリティーダービー', '').replace('ウマ娘プリティ', '')
    s = s.replace('賽馬娘PRETTYDERBY', '').replace('赛马娘PRETTYDERBY', '')
    s = s.replace('賽馬娘', '').replace('赛马娘', '').replace('ウマ娘', '')
    return s


def _bigrams(s):
    return set(s[i:i + 2] for i in range(len(s) - 1)) if len(s) > 1 else {s}


def _sim(a, b):
    if not a or not b:
        return 0.0
    if a == b:
        return 1.0
    A, B = _bigrams(a), _bigrams(b)
    inter = len(A & B)
    return 2.0 * inter / (len(A) + len(B)) if A and B else 0.0


def build_live_index():
    """title -> {url, date} from live_cat_data.json (mirrors SPA url scheme)."""
    try:
        data = json.load(io.open(LIVE_DATA, encoding='utf-8'))
    except Exception as ex:
        print('live index unavailable:', str(ex)[:80], flush=True)
        return []
    idx = []
    for cat, v in data.items():
        if not isinstance(v, dict):
            continue
        has_sections = bool(v.get('sections'))
        secs = v.get('sections') or [{'groups': v.get('groups') or []}]
        for si, sec in enumerate(secs):
            for gi, g in enumerate(sec.get('groups') or []):
                subs = g.get('subs') or []
                gkey = _norm(g.get('group'))
                # URL scheme mirrors SPA syncFromUrl:
                #   cd      -> /live/cd/{section}/{group}[/{perf}/0]
                #   others  -> /live/{cat}/{group}[/{perf}/0]
                base = '/zh-Hans/live/%s/%s%d' % (
                    cat, (str(si) + '/') if has_sections else '', gi)
                for pi, sub in enumerate(subs):
                    url = base
                    if pi > 0:
                        url += '/%d/0' % pi
                    t = sub.get('title') or ''
                    d = (sub.get('date') or '').split(' ')[0].split('（')[0]
                    idx.append({'key': _norm(t), 'raw': t, 'url': url, 'date': d,
                                'part': _part_mark(t)})
                if len(subs) <= 1:
                    idx.append({'key': gkey, 'raw': g.get('group'), 'url': base,
                                'date': '', 'part': ''})
    return idx


_PART_RE = _re.compile(r'DAY\d|[12]部|昼公演|夜公演')
_ANCILLARY_RE = _re.compile(r'閉会式|開会式|エンディングステージ|オープニングステージ|お渡し会|サイン会|特典[会會]')


def _part_mark(s):
    n = s or ''
    found = []
    for m in _PART_RE.finditer(n):
        t = m.group(0)
        if t not in found:
            found.append(t)
    return '+'.join(found)


_VOL_RE = _re.compile(r'VOL\.?\s*(\d+)')
_KAI_RE = _re.compile(r'第\s*(\d+)\s*回')


def _conflict(a, b):
    va, vb = set(_VOL_RE.findall(a)), set(_VOL_RE.findall(b))
    if va and vb and not (va & vb):
        return True
    ka, kb = set(_KAI_RE.findall(a)), set(_KAI_RE.findall(b))
    if ka and kb and not (ka & kb):
        return True
    return False


def build_series_index():
    """numbered-series (1st~7th EVENT) index from live_data.json/SERIES_GRID.
    One entry per DAY (deep link /{no}_EVENT/{perf}/{day}), dates inferred from
    the sub's start date because EN lists shows per day."""
    try:
        h = io.open(APP_JS, encoding='utf-8').read()
        LD = json.load(io.open(NUMBERED_LIVE_DATA, encoding='utf-8'))
        SG = json.loads(re.search(r'const SERIES_GRID = (\[.*?\]);\n', h, re.S).group(1))
    except Exception as ex:
        print('series index unavailable:', str(ex)[:80], flush=True)
        return []
    import datetime
    idx = []
    for gi, sg in enumerate(SG):
        if gi >= len(LD):
            break
        base = '/zh-Hans/live/number_series_event/%s_EVENT' % sg.get('no')
        for pi, sub in enumerate(LD[gi].get('subs') or []):
            days = sub.get('days') or [{}]
            raw_d = (sub.get('date') or '').split(' ')[0].replace('.', '-')
            parts = raw_d.split('-')
            if len(parts) > 3:
                parts = parts[:3]
            start = None
            try:
                start = datetime.date(*map(int, parts))
            except Exception:
                start = None
            for di in range(max(len(days), 1)):
                url = base if (pi == 0 and di == 0) else '%s/%d/%d' % (base, pi, di)
                day_obj = days[di] if di < len(days) else {}
                label = day_obj.get('label', '') if isinstance(day_obj, dict) else ''
                t = (sub.get('title') or '') + ((' ' + label) if label else '')
                d = ''
                if start is not None:
                    d = (start + datetime.timedelta(days=di)).isoformat()
                elif raw_d and '暂未公布' not in raw_d:
                    d = '-'.join(parts[:3])
                idx.append({'key': _norm(t), 'raw': t, 'url': url, 'date': d,
                            'part': _part_mark(t)})
    return idx


def attach_live(events, index):
    """Date-first matching: pin by date, confirm by title."""
    hit = miss = 0
    for ev in events:
        raw_t = ev['title'] or ''
        key = _norm(raw_t)
        if not key or not ev.get('date'):
            miss += 1
            continue
        if ev.get('live'):
            hit += 1
            continue
        # ancillary formats (ceremonies, handout/sign sessions, ending stages) have no setlist
        if _ANCILLARY_RE.search(raw_t):
            miss += 1
            continue
        part = _part_mark(raw_t + ' ' + (ev.get('times') or ''))
        yev = _YEAR_RE.search(raw_t)

        import datetime
        try:
            ev_d = datetime.date(*map(int, ev['date'].split('-'))).toordinal()
        except Exception:
            miss += 1
            continue

        def collect(day_window, min_sim):
            cands = []
            for e in index:
                if not e['key'] or not e['date']:
                    continue
                raw_l = e['raw'] or ''
                yl = _YEAR_RE.search(raw_l)
                if yev and yl and yev.group(1) != yl.group(1):
                    continue
                if _conflict(key, _norm(raw_l)) or _conflict(raw_t, raw_l):
                    continue
                try:
                    dd = abs(datetime.date(*map(int, e['date'].split('-'))).toordinal() - ev_d)
                except Exception:
                    continue
                if dd > day_window:
                    continue
                s = _sim(key, e['key'])
                if key in e['key'] or e['key'] in key:
                    s = max(s, 0.97)
                if s < min_sim:
                    continue
                if part and e['part'] and part != e['part']:
                    s -= 0.25
                cands.append((s, dd, e))
            return cands

        # pass A: same-day events, moderate title fit
        cands = collect(0, 0.40)
        # pass B: adjacent-day (Day1/Day2 style), stricter title fit
        if not cands:
            cands = collect(1, 0.60)
        if not cands:
            miss += 1
            continue
        # 「N部」: prefer candidates whose own title carries the same part marker
        mbu = _re.search(r'([12])部', raw_t + ' ' + (ev.get('times') or ''))
        if mbu and len(cands) > 1:
            want = mbu.group(1) + '部'
            pref = [c for c in cands if want in (c[2]['raw'] or '')]
            if pref:
                cands = pref
            else:
                want_idx = str(int(mbu.group(1)) - 1)
                same = [c for c in cands if c[2]['url'].endswith('/' + want_idx + '/0')]
                if same:
                    cands = same
        cands.sort(key=lambda x: (-x[0], x[1]))
        ev['live'] = cands[0][2]['url']
        hit += 1
    print('live links (date-first): matched=%d unmatched=%d' % (hit, miss), flush=True)


def _days(ds):
    import datetime
    try:
        return datetime.date(*map(int, ds.split('-'))).toordinal()
    except Exception:
        return 0


def load_existing():
    """Base dataset = previous events_data.json (never re-crawled)."""
    if os.path.exists(OUT):
        try:
            d = json.load(io.open(OUT, encoding='utf-8'))
            evs = d.get('events') or []
            out = {}
            for e in evs:
                k = e.get('link') or (e.get('date', '') + e.get('title', ''))
                if k:
                    out[k] = e
            return out
        except Exception as ex:
            print('existing data unreadable, full crawl:', str(ex)[:80], flush=True)
    return {}


def main():
    all_events = load_existing()
    base_n = len(all_events)
    print('base events: %d' % base_n, flush=True)
    fresh = []
    page = 1
    import datetime
    today = datetime.date.today().isoformat()
    while True:
        url = BASE + ('?page=%d&limit=%d' % (page, LIMIT)) if page > 1 else BASE + '?limit=%d' % LIMIT
        print('fetch page=%d ...' % page, flush=True)
        try:
            content = fetch(url)
        except Exception as ex:
            if page == 1 and base_n == 0:
                print('first page failed, abort:', str(ex)[:120], flush=True)
                return 1
            print('page failed (%s), stop here' % str(ex)[:80], flush=True)
            break
        blocks = extract_blocks(content)
        new_cnt = 0
        for b in blocks:
            ev = parse_block(b)
            if not ev:
                continue
            key = ev['link'] or (ev['date'] + ev['title'])
            if key in all_events:
                continue
            all_events[key] = ev
            fresh.append(ev)
            new_cnt += 1
        print('  blocks=%d new=%d total=%d' % (len(blocks), new_cnt, len(all_events)), flush=True)
        # actor page is newest-first: once a page yields nothing new, older pages won't either
        if len(blocks) < LIMIT or new_cnt == 0:
            break
        page += 1
        time.sleep(1.2)
    print('crawl done: +%d new' % (len(all_events) - base_n), flush=True)

    # export freshly crawled events into events_list.xlsx (insert at row 2)
    export_events_xlsx(fresh)

    # merge manually curated events (pre-linked to specific场次)
    man_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'manual_events.json')
    if os.path.exists(man_path):
        try:
            for row in json.load(io.open(man_path, encoding='utf-8')):
                k = row.get('link')
                if k and k not in all_events:
                    r = {kk: vv for kk, vv in row.items() if kk != '_manual'}
                    all_events[k] = r
            print('manual events merged:', flush=True)
        except Exception as ex:
            print('manual events load failed:', str(ex)[:80], flush=True)

    events = sorted(all_events.values(), key=lambda e: e['date'], reverse=True)
    index = build_live_index() + build_series_index()
    # match only rows still without a link: brand-new rows + old rows whose
    # counterpart场次 appeared / got its date announced since last run
    pending = [e for e in events if not e.get('live')]
    print('rows to match: %d' % len(pending), flush=True)
    attach_live(pending, index)
    linked_total = sum(1 for e in events if e.get('live'))
    data = {
        '_meta': {
            'source': BASE,
            'actor': 'ウマ娘 プリティーダービー(gal’up!)',
            'count': len(events),
            'scraped_at': today,
        },
        'events': events,
    }
    tmp = OUT + '.tmp'
    with io.open(tmp, 'w', encoding='utf-8') as f:
        json.dump(data, f, ensure_ascii=False, indent=1)
    os.replace(tmp, OUT)
    print('WROTE %s events=%d linked=%d scraped_at=%s' % (OUT, len(events), linked_total, today), flush=True)
    return 0


if __name__ == '__main__':
    sys.exit(main())
