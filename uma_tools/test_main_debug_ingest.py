"""Regression fixtures for recoverable source parsing and maintenance outcomes."""
import io
import json
import subprocess
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch
from PIL import Image
import anime_credit_text
import anime_web
import auto_albums
import crawl_events
import sync_anime_sources as sources
import sync_anime_assets as assets
import sync_anime_watch as watch
from video_fragment import select_video


class IngestRegressionTests(unittest.TestCase):
    def test_event_history_and_actor_pagination_survive_known_pages(self):
        import knowledge_io
        actor_page = 'https://www.eventernote.com/actors/Actor/1'
        old = {'link': 'https://www.eventernote.com/events/1', 'date': '2020-01-01', 'title': 'ウマ娘 old', 'actors': [{'name': 'Actor', 'url': actor_page}]}
        hole = {**old, 'link': 'https://www.eventernote.com/events/2', 'title': 'ウマ娘 historical hole'}
        personal = {**old, 'link': 'https://www.eventernote.com/events/3', 'title': 'ウマ娘 personal history'}
        state = {'attempted_at': {}}
        def read(file, default):
            return {'voice_actors': [{'identity': {'ja': 'Actor'}}]} if str(file).endswith('voice_actor_profiles.json') else state
        def blocks(url):
            if url.startswith(actor_page):
                return [old] * 100 if 'page=1' in url else [personal]
            return [hole] if 'page=3' in url else [old] * 100
        with tempfile.TemporaryDirectory() as directory:
            output = str(Path(directory) / 'events_data.json')
            Path(directory, 'manual_events.json').write_text('{broken', encoding='utf-8')
            with patch.object(crawl_events, '__file__', str(Path(directory) / 'crawler.py')), patch.object(crawl_events, 'OUT', output), \
                 patch.object(crawl_events, 'load_existing', return_value={old['link']: old}), patch.object(crawl_events, 'fetch', side_effect=lambda url: url), \
                 patch.object(crawl_events, 'extract_blocks', side_effect=blocks), patch.object(crawl_events, 'parse_block', side_effect=lambda row: dict(row)), \
                 patch.object(crawl_events, 'build_live_index', return_value=[]), patch.object(crawl_events, 'build_series_index', return_value=[]), \
                 patch.object(crawl_events, 'attach_live'), patch.object(crawl_events.time, 'sleep'), \
                 patch.object(knowledge_io, 'read_json', side_effect=read), patch.object(knowledge_io, 'write_json'), \
                 patch.object(crawl_events, 'export_events_xlsx', side_effect=lambda rows: self.assertTrue(Path(output).exists())):
                self.assertEqual(crawl_events.main(), 1, 'a malformed manual source reports failure')
            result = json.loads(Path(output).read_text())
            self.assertEqual({row['link'] for row in result['events']}, {old['link'], hole['link'], personal['link']})
            self.assertIn(str(Path(directory) / 'manual_events.json'), result['_meta']['source_failures'])

    def test_event_spreadsheet_export_is_idempotent(self):
        import openpyxl
        with tempfile.TemporaryDirectory() as directory:
            events = str(Path(directory) / 'events.xlsx')
            voices = str(Path(directory) / 'voices.xlsx')
            book = openpyxl.Workbook()
            book.active.append(['日期', '标题', '会场', '出演者'])
            book.save(events)
            book.active.append(['1', 'Actor'])
            book.save(voices)
            row = {'link': 'https://example.test/event/1', 'date': '2020-01-01', 'title': 'Event', 'venue': 'Venue', 'actors': [{'name': 'Actor'}]}
            with patch.object(crawl_events, 'EVENTS_XLSX', events), patch.object(crawl_events, 'VOICE_XLSX', voices):
                crawl_events.export_events_xlsx([row])
                crawl_events.export_events_xlsx([row])
            saved = openpyxl.load_workbook(events)
            self.assertEqual(saved.active.max_row, 2)
            self.assertEqual(saved.active.cell(2, 5).value, row['link'])
            saved.close()

    def test_episode_fallback_keeps_structured_credits(self):
        markup = '<script id="__NUXT_DATA__">bad json</script><h2 class="episode-title">第3話</h2><p class="story-text">Story</p>'
        fallback = anime_web.episode_records(markup, 'work', 'https://example.test/story/3')
        merged = sources.merge_episode_records([{'number': 3, 'cast': [{'name': 'Actor'}], 'summary_ja': 'Full story'}, *fallback])
        self.assertEqual(merged[0]['cast'], [{'name': 'Actor'}])
        self.assertEqual(merged[0]['summary_ja'], 'Full story')

    def test_column_boundary_and_song_title_do_not_create_false_credits(self):
        rows = anime_credit_text.staff_text_rows('監督\nPersonA\nキャスト\nPersonB')
        self.assertEqual(rows, [{'title': '監督', 'name': 'PersonA'}])
        self.assertEqual(anime_credit_text.music_rows('<ul class="music-list"><li>「STOP RED」</li></ul>'), [])
        rows = anime_credit_text.music_rows('<ul class="music-list"><li>OP 「STOP RED」</li></ul>')
        self.assertEqual(rows[0]['usage_type'], 'opening')

    def test_new_work_identity_includes_its_source(self):
        documents = {'works': {'works': []}, 'units': {'series': []}, 'source_registry': {'sources': []}, 'asset_sources': {'assets': []}}
        a = sources.discover_work({'url': 'https://a.test/s2/', 'title': 'A'}, documents, get=lambda _: '')
        b = sources.discover_work({'url': 'https://b.test/s2/', 'title': 'B'}, documents, get=lambda _: '')
        self.assertNotEqual(a['id'], b['id'])

    def test_widescreen_still_passes_after_resizing(self):
        body = io.BytesIO()
        Image.new('RGB', (1920, 800)).save(body, 'PNG')
        with tempfile.TemporaryDirectory() as directory, patch.object(assets, 'ROOT', directory):
            asset = {'output': 'still.webp', 'mode': 'still', 'source_url': 'https://example.test/still.png'}
            assets.write_cover(asset, fetch_image=lambda _: body.getvalue())
            assets.validate([asset])
            with Image.open(Path(directory) / 'still.webp') as image:
                self.assertEqual(image.size, (1280, 533))

    def test_membership_and_short_youtube_url(self):
        calls = []
        def run(args, **kwargs):
            calls.append(args)
            return subprocess.CompletedProcess(args, 0, json.dumps({'id': 'abcdefghijk', 'availability': 'subscriber_only'}), '')
        result = watch.check_resource({'url': 'https://youtu.be/abcdefghijk', 'platform': 'youtube'}, run=run)
        self.assertEqual(result[0], 'membership')
        self.assertIn('https://www.youtube.com/watch?v=abcdefghijk', calls[0])

    def test_provider_canonical_redirect_preserves_resource_identity(self):
        class Response(io.BytesIO):
            status = 200
            def geturl(self):
                return 'https://example.test/series/123/'
        with patch.object(watch.urllib.request, 'urlopen', return_value=Response('<title>赛马娘</title>'.encode())):
            self.assertEqual(watch.check_resource({'url': 'https://example.test/series/123', 'platform': 'other'})[0], 'active')

    def test_muxed_video_is_a_bounded_fallback(self):
        muxed = {'protocol': 'https', 'vcodec': 'h264', 'acodec': 'aac', 'height': 720, 'tbr': 600}
        video = {**muxed, 'acodec': 'none', 'height': 480}
        self.assertEqual(select_video({'formats': [muxed]}, 10), muxed)
        self.assertEqual(select_video({'formats': [muxed, video]}, 10), video)

    def test_album_outcomes_distinguish_failure_no_match_and_no_change(self):
        with patch.object(auto_albums, 'http_json', return_value=None):
            with self.assertRaises(RuntimeError):
                auto_albums.netease_search('Album')
        with patch.object(auto_albums, 'netease_search', return_value=(None, 0)), patch.object(auto_albums.time, 'sleep'):
            self.assertEqual(auto_albums.enrich_entry({'name': 'Album'}), 'no_match')
        entry = {'name': 'Album', 'count': '1 曲', 'cover': 'cover', 'neteaseAlbumId': '1', 'songs': [{'name': 'Song', 'url': 'manual'}]}
        with patch.object(auto_albums, 'netease_search', return_value=({'id': 1}, 100)), patch.object(auto_albums, 'netease_detail', return_value={'songs': [{'name': 'Song'}]}), patch.object(auto_albums, 'full_audio_ids', return_value=set()), patch.object(auto_albums.time, 'sleep'):
            self.assertEqual(auto_albums.enrich_entry(entry), 'unchanged')


if __name__ == '__main__':
    unittest.main()
